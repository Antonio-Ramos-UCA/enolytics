"""Plantilla visual de ENOLYTICS para presentaciones PPTX.

Reúne la identidad del cuadro de mando (albariza, escala de color del vino de Jerez,
tipografías) y los componentes que se repiten de una presentación a otra: láminas con
cabecera, separadores de sección, tarjetas, tablas, viñetas y estilo de gráficos nativos.

Uso:
    from plantilla_pptx import Baraja, tarjeta, vinetas, remate
    b = Baraja("ENOLYTICS · Universidad de Cádiz")
    s, y = b.lamina("Título", eyebrow="Sección", subtitulo="Bajada opcional")
    ...
    b.guardar("salida.pptx")

Nota: los gráficos se crean nativos (editables en PowerPoint), no como imagen.
Todas las coordenadas se fuerzan a entero: PowerPoint descarta la forma si el XML
lleva una coordenada fraccionaria, y lo hace en silencio.
"""
from __future__ import annotations

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_LEGEND_POSITION
from pptx.oxml.ns import qn
from lxml import etree

# --------------------------------------------------------------------------- #
# Identidad: cromo de la interfaz (escala del vino) + colores de DATOS validados
# --------------------------------------------------------------------------- #
ACENTO = RGBColor(0x7A, 0x3B, 0x1E)   # oloroso
INK = RGBColor(0x1C, 0x1B, 0x19)
INK2 = RGBColor(0x57, 0x54, 0x4E)
TENUE = RGBColor(0x8B, 0x87, 0x7E)
REGLA = RGBColor(0xD3, 0xCC, 0xBC)
PLANO = RGBColor(0xF3, 0xF0, 0xE9)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)
PX = RGBColor(0x3B, 0x1E, 0x12)
CREMA = RGBColor(0xF6, 0xEF, 0xE3)
FOCO_BG = RGBColor(0xFD, 0xF6, 0xF1)

AZUL = RGBColor(0x2A, 0x78, 0xD6)      # serie 1 (paleta validada del proyecto)
NARANJA = RGBColor(0xEB, 0x68, 0x34)   # serie 2
CONTEXTO = RGBColor(0xC9, 0xC5, 0xBA)  # «foco + contexto»
VERDE = RGBColor(0x0C, 0xA3, 0x0C)
AMBAR = RGBColor(0xC6, 0x90, 0x2F)

ESCALA = ["EBD98A", "C6902F", "A45E2A", "7A3B1E", "3B1E12"]  # fino → Pedro Ximénez

TITULO_F = "Georgia"
CUERPO_F = "Calibri"

W, H = Inches(13.333), Inches(7.5)
MX = Inches(0.75)
ANCHO = W - 2 * MX


def rgb(hexstr: str) -> RGBColor:
    return RGBColor.from_string(hexstr)


# --------------------------------------------------------------------------- #
# Primitivas
# --------------------------------------------------------------------------- #
def caja(slide, l, t, w, h, texto="", size=16, bold=False, color=INK, font=CUERPO_F,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(int(l), int(t), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = texto
    f = r.font
    f.size, f.bold, f.italic, f.name = Pt(size), bold, italic, font
    f.color.rgb = color
    return tb


def parrafo(tf, texto, size=16, color=INK2, bold=False, font=CUERPO_F, space_after=6,
            bullet=False, nivel=0, italic=False):
    """Añade un párrafo. Admite negritas delimitadas por **…**."""
    p = tf.add_paragraph() if (tf.paragraphs[0].runs or len(tf.paragraphs) > 1) else tf.paragraphs[0]
    p.space_after = Pt(space_after)
    if bullet:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(int(Inches(0.22 + 0.25 * nivel))))
        pPr.set("indent", str(-int(Inches(0.22))))
        bu = etree.SubElement(pPr, qn("a:buChar"))
        bu.set("char", "•" if nivel == 0 else "–")
    for i, trozo in enumerate(texto.split("**")):
        if not trozo:
            continue
        r = p.add_run()
        r.text = trozo
        f = r.font
        f.size, f.name, f.italic = Pt(size), font, italic
        f.bold = bold or (i % 2 == 1)
        f.color.rgb = INK if (i % 2 == 1 and color == INK2) else color
    return p


def rect(slide, l, t, w, h, fill, line=None):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(l), int(t), int(w), int(h))
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(0.75)
    s.shadow.inherit = False
    return s


def cuerpo(slide, l, t, w, h):
    """Marco de texto vacío, listo para `parrafo()`."""
    tb = slide.shapes.add_textbox(int(l), int(t), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Emu(0)
    tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


# --------------------------------------------------------------------------- #
# Componentes
# --------------------------------------------------------------------------- #
def vinetas(slide, l, t, w, h, items, size=16, space=7):
    tf = cuerpo(slide, l, t, w, h)
    for it in items:
        if isinstance(it, tuple):      # (texto,) en segundo nivel
            parrafo(tf, it[0], size=size - 1, bullet=True, nivel=1, space_after=space - 2)
        else:
            parrafo(tf, it, size=size, bullet=True, space_after=space)
    return tf


def tarjeta(slide, l, t, w, h, rotulo, valor=None, texto=None, foco=False,
            size_val=30, size_txt=12.5):
    rect(slide, l, t, w, h, BLANCO, line=REGLA)
    if foco:
        rect(slide, l, t, Inches(0.06), h, ACENTO)
    x = l + Inches(0.22)
    caja(slide, x, t + Inches(0.16), w - Inches(0.4), Inches(0.3), rotulo.upper(),
         size=9, bold=True, color=TENUE)
    y = t + Inches(0.46)
    if valor:
        caja(slide, x, y, w - Inches(0.4), Inches(0.6), valor, size=size_val, bold=True,
             color=INK, font=TITULO_F)
        y += Inches(0.62)
    if texto:
        tf = cuerpo(slide, x, y, w - Inches(0.4), h - (y - t) - Inches(0.15))
        parrafo(tf, texto, size=size_txt, color=INK2, space_after=0)


def panel(slide, l, t, w, h, rotulo, parrafos, foco=True, size=12):
    """Bloque destacado: rótulo en acento y varios párrafos."""
    rect(slide, l, t, w, h, FOCO_BG if foco else PLANO, line=ACENTO if foco else None)
    caja(slide, l + Inches(0.25), t + Inches(0.15), w - Inches(0.5), Inches(0.3),
         rotulo.upper(), size=10, bold=True, color=ACENTO)
    tf = cuerpo(slide, l + Inches(0.25), t + Inches(0.47), w - Inches(0.5), h - Inches(0.62))
    for i, p_ in enumerate(parrafos):
        parrafo(tf, p_, size=size, space_after=6 if i < len(parrafos) - 1 else 0)


def tabla(slide, l, t, w, filas, anchos, size=12, cab_size=10, alto_fila=Inches(0.42)):
    n_r, n_c = len(filas), len(filas[0])
    shp = slide.shapes.add_table(n_r, n_c, int(l), int(t), int(w), int(alto_fila * n_r))
    tb = shp.table
    for j, a in enumerate(anchos):
        tb.columns[j].width = int(a)
    for i, fila in enumerate(filas):
        for j, val in enumerate(fila):
            c = tb.cell(i, j)
            c.margin_left = c.margin_right = Inches(0.08)
            c.margin_top = c.margin_bottom = Inches(0.04)
            c.fill.solid()
            c.fill.fore_color.rgb = PLANO if i == 0 else (BLANCO if i % 2 else rgb("FBFAF7"))
            p = c.text_frame.paragraphs[0]
            c.text_frame.word_wrap = True
            for k, trozo in enumerate(str(val).split("**")):
                if not trozo:
                    continue
                r = p.add_run()
                r.text = trozo
                r.font.size = Pt(cab_size if i == 0 else size)
                r.font.name = CUERPO_F
                r.font.bold = (i == 0) or (k % 2 == 1)
                r.font.color.rgb = TENUE if i == 0 else INK
    return tb


def remate(slide, texto, y=None, size=17):
    y = y if y is not None else H - Inches(1.45)
    rect(slide, MX, y, ANCHO, Emu(19050), ACENTO)
    caja(slide, MX, y + Inches(0.12), ANCHO, Inches(0.6), texto, size=size, bold=True,
         color=INK, font=TITULO_F)


def cita_pie(slide, texto):
    caja(slide, MX, H - Inches(0.78), ANCHO, Inches(0.3), texto, size=9.5,
         color=TENUE, italic=True)


def estilo_chart(chart, leyenda=False):
    chart.font.name = CUERPO_F
    chart.font.size = Pt(11)
    chart.font.color.rgb = INK2
    chart.has_legend = leyenda
    if leyenda:
        chart.legend.position = XL_LEGEND_POSITION.TOP
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(11)
    for ax in (chart.category_axis, chart.value_axis):
        ax.format.line.color.rgb = REGLA
        ax.tick_labels.font.size = Pt(11)
        ax.tick_labels.font.color.rgb = TENUE
    chart.value_axis.major_gridlines.format.line.color.rgb = rgb("E4E0D6")
    chart.category_axis.has_major_gridlines = False


def titulo_chart(chart, texto):
    chart.has_title = True
    chart.chart_title.text_frame.text = texto
    r = chart.chart_title.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.color.rgb, r.font.name = Pt(11), True, INK, CUERPO_F


# --------------------------------------------------------------------------- #
# La baraja
# --------------------------------------------------------------------------- #
class Baraja:
    """Presentación 16:9 con la identidad de ENOLYTICS."""

    def __init__(self, pie: str):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self._blank = self.prs.slide_layouts[6]
        self._n = 0
        self.pie_texto = pie

    # --- piezas comunes ---
    def _filo(self, slide, t=Emu(0), h=Inches(0.07)):
        seg = W / len(ESCALA)
        for i, c in enumerate(ESCALA):
            rect(slide, int(seg * i), t, int(seg) + 1, h, rgb(c))

    def _pie(self, slide, oscuro=False):
        self._n += 1
        col = CREMA if oscuro else TENUE
        caja(slide, MX, H - Inches(0.42), Inches(8.5), Inches(0.3), self.pie_texto,
             size=9, color=col)
        caja(slide, W - MX - Inches(1), H - Inches(0.42), Inches(1), Inches(0.3),
             str(self._n), size=9, color=col, align=PP_ALIGN.RIGHT)

    # --- tipos de lámina ---
    def lamina(self, titulo, eyebrow=None, subtitulo=None):
        """Lámina estándar. Devuelve (slide, y) con y = primera coordenada libre."""
        s = self.prs.slides.add_slide(self._blank)
        self._filo(s)
        y = Inches(0.45)
        if eyebrow:
            caja(s, MX, y, ANCHO, Inches(0.3), eyebrow.upper(), size=10, bold=True, color=ACENTO)
            y += Inches(0.32)
        caja(s, MX, y, ANCHO, Inches(0.7), titulo, size=30 if len(titulo) <= 52 else 25,
             bold=True, color=INK, font=TITULO_F)
        y += Inches(0.72)
        if subtitulo:
            caja(s, MX, y, ANCHO, Inches(0.55), subtitulo, size=15, color=INK2)
            y += Inches(0.6)
        rect(s, MX, y, ANCHO, Emu(9525), REGLA)
        y += Inches(0.22)
        self._pie(s)
        return s, y

    def oscura(self, titulo, eyebrow=None, bajada=None, numero=None):
        """Separador o cierre sobre fondo Pedro Ximénez."""
        s = self.prs.slides.add_slide(self._blank)
        rect(s, 0, 0, W, H, PX)
        self._filo(s, t=H - Inches(0.07))
        y = Inches(2.0)
        if numero:
            caja(s, MX, Inches(1.5), ANCHO, Inches(0.6), numero, size=40, bold=True,
                 color=rgb("5A3A28"), font=CUERPO_F)
        if eyebrow:
            caja(s, MX, y, ANCHO, Inches(0.3), eyebrow.upper(), size=11, bold=True,
                 color=rgb("EBD98A"))
            y += Inches(0.42)
        caja(s, MX, y, Inches(11.3), Inches(1.5), titulo, size=42, bold=True,
             color=CREMA, font=TITULO_F)
        if bajada:
            caja(s, MX, y + Inches(1.6), Inches(10), Inches(1.2), bajada, size=17,
                 color=rgb("C9BEAC"))
        self._pie(s, oscuro=True)
        return s

    def guardar(self, ruta):
        import pathlib
        p = pathlib.Path(ruta)
        p.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(p))
        return p, len(self.prs.slides.__iter__.__self__._sldIdLst)


def notas(slide, texto):
    slide.notes_slide.notes_text_frame.text = texto.strip()
