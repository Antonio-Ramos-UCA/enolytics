"""Presentación breve de ENOLYTICS para GESTORES DE DESTINO.

Audiencia: patronatos de turismo, rutas del vino, OGD y ayuntamientos interesados en
conocer la herramienta. El hilo no es el método sino la decisión: qué le dice la
plataforma a quien gestiona un destino, qué recibe su tejido de bodegas y qué haría
falta para llevarla a otro territorio. Una lámina con la ficha técnica.

Cifras verificadas contra datos/procesado. Uso:
    python3 scripts/generar_presentacion_gestores.py [salida.pptx]
"""
import sys
import pathlib

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))

from pptx.util import Inches, Pt
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN

from plantilla_pptx import (
    Baraja, caja, rect, cuerpo, parrafo, vinetas, tarjeta, panel, tabla,
    remate, cita_pie, estilo_chart, titulo_chart, notas, rgb,
    ACENTO, INK, INK2, TENUE, REGLA, PLANO, BLANCO, CREMA, FOCO_BG,
    AZUL, NARANJA, CONTEXTO, VERDE, MX, ANCHO, W, H, TITULO_F,
)

b = Baraja("ENOLYTICS · Universidad de Cádiz · proyecto FEDER Andalucía 2021-2027")

# ════════════════════════════════ 1 · PORTADA ════════════════════════════════
s = b.oscura(
    "Inteligencia competitiva para la gestión de un destino enoturístico",
    eyebrow="Proyecto ENOLYTICS · Marco de Jerez",
    bajada="Una plataforma que convierte datos públicos —reseñas, estadística oficial, webs— en "
           "decisiones para el destino y para cada bodega. Sin pedir un solo dato a las empresas.")
caja(s, MX, Inches(5.9), Inches(11), Inches(0.9),
     "Universidad de Cádiz · Instituto INDESS\n"
     "IP: María Paula Lechuga Sancho · Co-IP: Antonio Rafael Ramos Rodríguez",
     size=12, color=rgb("C9BEAC"))
notas(s, """
Audiencia: gestores de destino. En una frase, lo que se viene a contar: hemos construido una
herramienta que ya funciona sobre el Marco de Jerez, y lo que hace se puede replicar en otro
destino.
Subrayar desde el principio el argumento que más les va a interesar: NO hemos pedido datos a las
bodegas. Todo sale de fuentes públicas. Eso elimina la barrera habitual de estos proyectos.
""")

# ════════════════════════════ 2 · EL PROBLEMA ════════════════════════════
s, y = b.lamina("Con qué información se decide hoy en un destino",
                eyebrow="El punto de partida")
cw = (ANCHO - Inches(0.5)) / 3
tarjeta(s, MX, y, cw, Inches(1.9), "Lo que sí hay", None,
        "Estadística agregada anual: visitantes, ingresos, perfil de la demanda. Llega **tarde** "
        "y **no baja** al nivel de decisión: qué falla, en qué bodega, en qué mes.")
tarjeta(s, MX + cw + Inches(0.25), y, cw, Inches(1.9), "Lo que falta", None,
        "Qué opina realmente el visitante, atributo por atributo. Cómo va el destino **frente a "
        "otras rutas**. Qué oportunidades hay fuera del radar habitual.")
tarjeta(s, MX + 2 * (cw + Inches(0.25)), y, cw, Inches(1.9), "Y el dato existe", None,
        "Hay **decenas de miles de opiniones públicas** sobre las bodegas del destino, y "
        "estadística oficial abierta que nadie estaba cruzando entre sí.", foco=True)
vinetas(s, MX, y + Inches(2.2), ANCHO, Inches(2.0), [
    "El sector enoturístico crece —**2,98 millones** de visitantes a bodegas y museos en España, "
    "**+18,2 %** en un año— y con el crecimiento llega la necesidad de diferenciarse.",
    "El Marco de Jerez está entre las **tres rutas más visitadas** de España: **425.652 visitantes** "
    "y **17,2 M€** de ingresos en 2024.",
    "Pero crecer en volumen no es lo mismo que **captar valor**: el Marco ingresa **40,4 € por "
    "visitante**, frente a los **70,0 €** de la ruta líder.",
], size=15, space=8)
cita_pie(s, "Fuentes: ACEVIN, Observatorio Turístico de las Rutas del Vino de España (2023, 2024).")
notas(s, """
Diagnóstico sin reproche: no es que los gestores decidan mal, es que no había otra información
disponible. La estadística oficial es imprescindible pero llega agregada y con retraso.
El dato de los 40,4 € frente a 70,0 € suele ser el que capta la atención: es una brecha de
monetización, no de volumen. Guardarlo porque reaparece al final.
""")

# ════════════════════════════ 3 · QUÉ ES ════════════════════════════
s, y = b.lamina("Qué es ENOLYTICS", eyebrow="La herramienta",
                subtitulo="Un cuadro de mando web que ya está funcionando sobre las 42 bodegas del Marco de Jerez.")
cw2 = (ANCHO - Inches(0.5)) / 2
tarjeta(s, MX, y, cw2, Inches(2.0), "Para quien gestiona el destino", None,
        "Cómo va el destino en **siete dimensiones competitivas**, dónde están las fugas de valor y "
        "qué acciones priorizar. Con un semáforo de situación y **14 recomendaciones** ordenadas por "
        "urgencia, cada una con el dato que la justifica.", foco=True)
tarjeta(s, MX + cw2 + Inches(0.5), y, cw2, Inches(2.0), "Para cada empresa del destino", None,
        "Su reputación en detalle, **quiénes son sus competidores directos** —no las 41 bodegas, "
        "sino tres o cuatro—, sus puntos fuertes y débiles, y un **informe de una página** "
        "descargable, pensado para entregárselo.")
vinetas(s, MX, y + Inches(2.3), ANCHO, Inches(1.9), [
    "**Está vivo, no es un informe.** Se abre en el navegador y los datos se actualizan; no es una "
    "foto de un mes concreto.",
    "**No pide nada a las empresas.** Se alimenta de fuentes públicas: eso permite tener el análisis "
    "del destino completo desde el primer día, sin esperar a que nadie rellene un cuestionario.",
    "**Dice lo que no sabe.** Cuando no hay muestra suficiente para afirmar algo, lo declara en vez "
    "de rellenar el hueco.",
], size=15, space=9)
remate(s, "Está publicado y es accesible hoy: enolytics-feder-uca.streamlit.app")
notas(s, """
Las tres viñetas son las tres objeciones habituales, contestadas antes de que se formulen:
«¿esto se queda obsoleto?», «¿tengo que convencer a las bodegas para que aporten datos?» y
«¿me fío de lo que dice?».
Ofrecer abrirlo en directo al final si hay tiempo y conexión.
""")

# ════════════════════════════ 4 · LAS SIETE ════════════════════════════
s, y = b.lamina("Siete dimensiones, un solo cuadro de mando", eyebrow="El modelo de análisis",
                subtitulo="El marco de inteligencia competitiva del proyecto, alineado con el modelo de Destinos Turísticos Inteligentes de SEGITTUR y la norma UNE 178502.")
filas = [
    ["Dimensión", "Qué responde para el destino", "Ejemplo de lo que ya calcula"],
    ["Económica", "¿Cuánto valor capta el destino por visitante?", "Ingreso por visitante y brecha con otras rutas"],
    ["Mercado", "¿Quién nos busca y desde dónde?", "Interés de búsqueda y demanda aérea no atendida"],
    ["Competidores", "¿Cómo estamos frente a otros destinos?", "Posición en el ranking de rutas y reputación comparada"],
    ["Clientes", "¿Qué valora y qué critica el visitante?", "Satisfacción por atributo sobre 10.972 opiniones"],
    ["Negocios", "¿Qué oferta y qué capacidad hay?", "Censo de bodegas, servicios y oferta enoturística"],
    ["Tecnológica", "¿Qué madurez digital y accesibilidad tiene la oferta?", "Índice digital por bodega y barreras de accesibilidad"],
    ["**Sostenibilidad**", "**¿Cómo de sostenible es la actividad?**", "**Certificación ambiental y acceso en transporte público**"],
]
tabla(s, MX, y, ANCHO, filas, [Inches(1.9), Inches(4.9), Inches(5.03)], size=11.5, alto_fila=Inches(0.42))
caja(s, MX, y + Inches(3.5), ANCHO, Inches(0.5),
     "La dimensión de sostenibilidad es la aportación propia del proyecto: tratarla como una fuente "
     "de inteligencia más, no como un apartado del informe anual.", size=13, color=INK2)
notas(s, """
No leer la tabla entera: señalar dos filas. Recomiendo «Clientes» (porque es la que sorprende, por
el volumen de opiniones) y «Sostenibilidad» (porque es la aportación del proyecto y porque conecta
con las obligaciones que ya tienen los destinos).
Si preguntan por SEGITTUR: los indicadores se organizan en oferta, demanda y sostenibilidad, las
tres categorías del modelo DTI. La UNE 178502 es la norma que lo formaliza.
""")

# ════════════════════════════ 5 · LOS DATOS ════════════════════════════
s, y = b.lamina("De dónde sale cada dato", eyebrow="Fuentes y transparencia",
                subtitulo="Cada cifra de la plataforma lleva una marca que dice cómo de firme es. Nunca se mezcla estadística oficial con una estimación.")
cw = (ANCHO - Inches(0.5)) / 3
rect(s, MX, y, cw, Inches(2.5), BLANCO, line=REGLA)
rect(s, MX, y, Inches(0.06), Inches(2.5), VERDE)
caja(s, MX + Inches(0.25), y + Inches(0.16), cw - Inches(0.5), Inches(0.3), "DATO OFICIAL", size=10, bold=True, color=VERDE)
tf = cuerpo(s, MX + Inches(0.25), y + Inches(0.5), cw - Inches(0.5), Inches(1.9))
parrafo(tf, "**ACEVIN** — visitantes, ingresos, oferta y perfil de la demanda por ruta.", size=12.5, space_after=6)
parrafo(tf, "**Dataestur / SEGITTUR** — gasto turístico, percepción, cruceros y conectividad aérea.", size=12.5, space_after=6)
parrafo(tf, "**FEV** — certificación ambiental de bodegas.", size=12.5, space_after=6)
parrafo(tf, "**Ruta del Vino y Brandy** — catálogo de las 42 bodegas.", size=12.5, space_after=0)

x2 = MX + cw + Inches(0.25)
rect(s, x2, y, cw, Inches(2.5), BLANCO, line=REGLA)
rect(s, x2, y, Inches(0.06), Inches(2.5), AZUL)
caja(s, x2 + Inches(0.25), y + Inches(0.16), cw - Inches(0.5), Inches(0.3), "DATO OBSERVADO", size=10, bold=True, color=AZUL)
tf = cuerpo(s, x2 + Inches(0.25), y + Inches(0.5), cw - Inches(0.5), Inches(1.9))
parrafo(tf, "**10.972 reseñas** de Google sobre las bodegas del Marco, de 2011 a 2026, en seis idiomas.", size=12.5, space_after=6)
parrafo(tf, "**Google Trends** — interés de búsqueda frente a Rioja, Ribera y Rías Baixas.", size=12.5, space_after=6)
parrafo(tf, "**Auditoría propia** de las webs de las bodegas y **OpenStreetMap** para el transporte.", size=12.5, space_after=0)

x3 = MX + 2 * (cw + Inches(0.25))
rect(s, x3, y, cw, Inches(2.5), BLANCO, line=REGLA)
rect(s, x3, y, Inches(0.06), Inches(2.5), rgb("C6902F"))
caja(s, x3 + Inches(0.25), y + Inches(0.16), cw - Inches(0.5), Inches(0.3), "ESTIMACIÓN", size=10, bold=True, color=rgb("C6902F"))
tf = cuerpo(s, x3 + Inches(0.25), y + Inches(0.5), cw - Inches(0.5), Inches(1.9))
parrafo(tf, "**Sin fuente propia.** Es un cálculo nuestro a partir de los anteriores: ingreso por "
            "visitante, estacionalidad aproximada, procedencia deducida del idioma.", size=12.5, space_after=6)
parrafo(tf, "Se marca precisamente **para que se pueda dudar de ello**.", size=12.5, space_after=0)

remate(s, "Ninguna bodega ha tenido que aportar un solo dato. Todo es información pública, tratada de forma agregada.",
       y=y + Inches(2.8))
cita_pie(s, "Los textos de la plataforma no los genera una inteligencia artificial: se componen de los datos agregados, para que el mismo dato produzca siempre el mismo texto.")
notas(s, """
Esta lámina compra la credibilidad del resto. Dedicarle tiempo.
El ámbar es el punto fuerte del discurso: reconocer lo que es cálculo propio distingue un análisis
de un folleto.
Si preguntan por protección de datos: son reseñas públicas, el nombre del autor no interviene en
ningún análisis y se trabaja con textos agregados por bodega y atributo, nunca con personas.
""")

# ════════════════════════════ 6 · LO QUE VE EL GESTOR ════════════════════════════
s, y = b.lamina("Lo que ve quien gestiona el destino", eyebrow="La vista de destino",
                subtitulo="Entrar, mirar cinco segundos y saber por dónde empezar.")
cw2 = (ANCHO - Inches(0.5)) / 2
vinetas(s, MX, y, cw2, Inches(2.6), [
    "**Las cifras del destino** de un vistazo: visitantes, bodegas, valoración media e ingresos.",
    "**Un semáforo de las siete dimensiones**, deducido de los avisos que genera el motor: qué está "
    "crítico, qué requiere atención y qué está en orden.",
    "**Las tres acciones más urgentes**, con su diagnóstico y su fuente. La lista completa, al lado.",
    "**El detalle por dimensión**, con las recomendaciones de cada una junto al dato que las justifica.",
], size=14, space=9)
panel(s, MX + cw2 + Inches(0.5), y, cw2, Inches(2.6), "Ejemplo real de recomendación", [
    "**Responder a las críticas sin contestar.**",
    "*Diagnóstico:* 590 reseñas de 1-2 estrellas sin una sola respuesta en el destino; 11 de las 33 "
    "bodegas con muestra suficiente no han contestado nunca.",
    "*Acción:* establecer una pauta de respuesta. Es gratis, lo lee quien está decidiendo la visita "
    "y no requiere inversión.",
], size=12.5)
rect(s, MX, y + Inches(2.9), ANCHO, Inches(1.3), PLANO)
caja(s, MX + Inches(0.3), y + Inches(3.05), ANCHO - Inches(0.6), Inches(0.3),
     "LOS UMBRALES ESTÁN DECLARADOS", size=10, bold=True, color=ACENTO)
caja(s, MX + Inches(0.3), y + Inches(3.35), ANCHO - Inches(0.6), Inches(0.75),
     "Cada recomendación salta cuando un dato cruza un listón, y la plataforma publica cuál es ese "
     "listón y de dónde sale: unos tienen base estadística, otros son criterio del equipo y se dicen "
     "como tales. Un sistema de recomendaciones con los umbrales a la vista es auditable.",
     size=12.5, color=INK2)
notas(s, """
El ejemplo de las reseñas sin responder funciona muy bien con gestores: es una acción de coste cero
sobre la que pueden actuar mañana, y es medible.
Si preguntan: la plataforma genera 14 recomendaciones para el destino y recomendaciones específicas
para cada bodega.
""")

# ════════════════════════════ 7 · CRUCEROS ════════════════════════════
s, y = b.lamina("Lo que aparece al cruzar fuentes que nadie cruzaba", eyebrow="Un hallazgo de destino",
                subtitulo="En 2024 desembarcaron 696.151 cruceristas en la Bahía de Cádiz. Todas las bodegas del Marco juntas recibieron 425.652 visitantes.")
cd = CategoryChartData()
cd.categories = ["Ene", "Feb", "Mar", "Abr", "May", "Jun", "Jul", "Ago", "Sep", "Oct", "Nov", "Dic"]
cd.add_series("Cruceristas (índice, 100 = su mes máximo)", (23, 9, 29, 66, 63, 41, 70, 68, 84, 100, 85, 39))
cd.add_series("Afluencia enoturística (índice; estimación)", (38, 44, 57, 60, 62, 60, 68, 100, 77, 75, 54, 51))
gf = s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, int(MX), int(y), int(Inches(8.2)), int(Inches(3.7)), cd)
ch = gf.chart
estilo_chart(ch, leyenda=True)
pl = ch.plots[0]
for i, col in enumerate((NARANJA, AZUL)):
    sr = pl.series[i]
    sr.format.line.color.rgb = col
    sr.format.line.width = Pt(2.5)
    sr.smooth = False
    sr.marker.size = 7
    sr.marker.format.fill.solid()
    sr.marker.format.fill.fore_color.rgb = col
    sr.marker.format.line.color.rgb = BLANCO
ch.value_axis.minimum_scale = 0
ch.value_axis.maximum_scale = 100

xr = MX + Inches(8.5)
wr = ANCHO - Inches(8.5)
tarjeta(s, xr, y, wr, Inches(1.5), "En noviembre", "85 vs 54",
        "Los cruceros están al 85 % de su máximo; el enoturismo, al 54 %.", foco=True, size_val=26)
tf = cuerpo(s, xr, y + Inches(1.7), wr, Inches(2.0))
parrafo(tf, "Cádiz es el **6.º puerto de España** en cruceros, y atraca a **40 minutos** de Jerez.", size=12.5, space_after=7)
parrafo(tf, "En torno al **30 % de los enoturistas no pernocta**: es el perfil de quien baja de un "
            "barco por unas horas.", size=12.5, space_after=7)
parrafo(tf, "**Lo accionable:** un producto de medio día empaquetado con navieras y agencias "
            "receptivas, para octubre y noviembre. Hoy no existe.", size=12.5, space_after=0)
cita_pie(s, "Cruceros: Dataestur / Puertos del Estado (oficial). La curva de afluencia enoturística es una estimación propia por volumen mensual de reseñas: sirve para la forma, no para dar cifras.")
notas(s, """
Declarar que la curva azul es una estimación antes de que lo pregunten: no existe afluencia mensual
publicada por bodega.
El mensaje de fondo: este hallazgo no está en ninguna reseña ni en ningún informe. Sale de poner dos
ficheros públicos uno al lado del otro. Eso es lo que hace la plataforma de forma sistemática.
Para un gestor, además, conecta con la prioridad de desestacionalización que ya tiene en su agenda.
""")

# ════════════════════════════ 8 · CONECTIVIDAD ════════════════════════════
s, y = b.lamina("Demanda que existe y no se está atendiendo", eyebrow="Un hallazgo de mercado")
cw2 = (ANCHO - Inches(0.5)) / 2
cd = CategoryChartData()
cd.categories = ["Italia", "Estados Unidos", "Francia", "Países Bajos"]
cd.add_series("Búsquedas de vuelo", (633809, 621324, 590577, 421559))
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, int(MX), int(y), int(cw2), int(Inches(2.6)), cd)
ch = gf.chart
estilo_chart(ch)
titulo_chart(ch, "Búsquedas de vuelo a Jerez desde países sin un solo asiento directo")
pl = ch.plots[0]
pl.gap_width = 55
pl.has_data_labels = True
pl.data_labels.font.size = Pt(10)
pl.data_labels.number_format = '#,##0'
pl.data_labels.number_format_is_linked = False
pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
ser = pl.series[0]
ser.format.fill.solid()
ser.format.fill.fore_color.rgb = CONTEXTO
pt = ser.points[1]
pt.format.fill.solid()
pt.format.fill.fore_color.rgb = AZUL
ch.value_axis.visible = False
ch.value_axis.has_major_gridlines = False

xr = MX + cw2 + Inches(0.5)
tarjeta(s, xr, y, cw2, Inches(1.45), "Total desde países sin vuelo directo", "3.253.175",
        "búsquedas. Solo **ocho países extranjeros** tienen asiento directo a Jerez.", foco=True, size_val=26)
panel(s, xr, y + Inches(1.62), cw2, Inches(1.0), "El caso de Estados Unidos",
      ["Es el mercado que **mejor convierte** la búsqueda en reserva (5,87 %) y el que **más noches se "
       "queda** (12,8 días). Y no tiene vuelo directo."], size=12.5)
vinetas(s, MX, y + Inches(2.85), ANCHO, Inches(1.4), [
    "**Es un indicador que mira al futuro:** las búsquedas de vuelo se hacen meses antes del viaje, "
    "así que informan de demanda que aún no ha ocurrido. Casi todo lo demás en turismo mira al pasado.",
    "**Y dice cuándo actuar:** Alemania planifica con 167 días de antelación; Estados Unidos, con 120. "
    "Eso fija la fecha de lanzamiento de cada campaña.",
], size=13.5, space=7)
cita_pie(s, "Fuente: Dataestur — búsquedas, reservas y asientos programados por país emisor, con datos hasta agosto de 2026.")
notas(s, """
Este es el dato con más valor político para un gestor: convierte «creemos que hay demanda» en «hay
3,25 millones de búsquedas y cero asientos», que es lo que se puede llevar a una mesa de negociación
de rutas aéreas con una aerolínea o con el aeropuerto.
Ofrecerles los datos: están por país, por mes y con antelación de compra.
""")

# ════════════════════════════ 9 · LO QUE RECIBE LA BODEGA ════════════════════════════
s, y = b.lamina("Y lo que puede entregar a cada empresa de su destino", eyebrow="La vista de bodega",
                subtitulo="El mismo análisis, bajado al nivel de cada establecimiento. Es lo que convierte la plataforma en un servicio al tejido empresarial.")
cw = (ANCHO - Inches(0.5)) / 3
tarjeta(s, MX, y, cw, Inches(2.2), "Su reputación en detalle", None,
        "Valoración, evolución, distribución de estrellas y un indicador que no muestra ninguna "
        "plataforma comercial: **la tasa de respuesta** del propietario a las reseñas y a las críticas.")
tarjeta(s, MX + cw + Inches(0.25), y, cw, Inches(2.2), "Sus competidores directos", None,
        "No las 41 bodegas restantes: **tres o cuatro**, identificadas por parecido en el cliente al "
        "que sirven y en los recursos de que disponen. Con un mapa que sitúa a cada rival.", foco=True)
tarjeta(s, MX + 2 * (cw + Inches(0.25)), y, cw, Inches(2.2), "Un informe de una página", None,
        "Descargable en PDF: reputación, competidores, fortalezas y debilidades, presencia digital y "
        "sostenibilidad, y qué hacer primero. **Pensado para entregárselo a la bodega.**")
vinetas(s, MX, y + Inches(2.5), ANCHO, Inches(1.7), [
    "**37 de las 42 bodegas** obtienen análisis de competidores. A las cinco restantes la plataforma "
    "les dice que no hay muestra suficiente, en lugar de darles competidores al azar.",
    "Para un gestor de destino esto es **un servicio que puede ofrecer a sus asociados** sin pedirles "
    "nada a cambio y sin coste por empresa.",
], size=14, space=8)
notas(s, """
Para un patronato o una ruta del vino, esta lámina responde a «¿y esto qué les doy yo a mis
empresas?». El informe de una página es tangible: se descarga y se entrega.
Si hay tiempo en la demo, descargar uno en directo. Es el momento que más engancha.
""")

# ════════════════════════════ 10 · FICHA TÉCNICA ════════════════════════════
s, y = b.lamina("Ficha técnica", eyebrow="Cómo está construida · para consultar con el área de informática")
cw2 = (ANCHO - Inches(0.5)) / 2
filas = [
    ["Aspecto", "Detalle"],
    ["Tecnología", "Python 3.10 y Streamlit. 25 módulos, unas 8.400 líneas. Sin licencias de software"],
    ["Arquitectura", "Ingesta → ETL → procesamiento de lenguaje → analítica → cuadro de mando"],
    ["Fuentes conectadas", "13 módulos de ingesta; 19 conjuntos de datos procesados que alimentan la aplicación"],
    ["Análisis de texto", "Léxico multilingüe propio (ES, EN, DE, FR, IT) + modelo BERT multilingüe de sentimiento"],
    ["Volumen analizado", "10.972 reseñas · 7.258 con texto · 17.246 menciones anotadas por atributo"],
    ["Modelos analíticos", "Importancia-desempeño (IPA, IPCA, DIPA), importancia por impacto (PRCA), perfil de Kano, competidores, co-ocurrencias"],
    ["Despliegue", "Servicio en la nube, acceso por navegador, sin instalación. Código abierto"],
    ["Coste recurrente", "Decenas de euros al año: actualizar el corpus de reseñas cuesta unos 36 $"],
    ["Verificación", "Pruebas automáticas de las cinco vistas y revisión en navegador real tras cada cambio"],
]
tabla(s, MX, y, ANCHO, filas, [Inches(2.4), Inches(9.43)], size=11, cab_size=10, alto_fila=Inches(0.35))
panel(s, MX, y + Inches(3.92), ANCHO, Inches(1.05), "Una decisión de diseño que conviene conocer",
      ["Los textos de la plataforma **no los redacta una inteligencia artificial generativa**: se componen "
       "a partir de los datos agregados. El mismo dato produce siempre el mismo texto, y eso hace que "
       "todo resultado sea reproducible y auditable."], size=12)
notas(s, """
Lámina de apoyo: no leerla entera. Señalar tres cosas y pasar:
1) Sin licencias ni instalación: se abre en un navegador.
2) El coste recurrente es ridículo comparado con cualquier plataforma comercial de inteligencia
   turística.
3) No hay IA generativa escribiendo conclusiones, que es la pregunta de 2026.
Si hay un técnico en la sala, ofrecerle el repositorio: el código es abierto.
""")

# ════════════════════════════ 11 · LLEVARLO A OTRO DESTINO ════════════════════════════
s, y = b.lamina("Qué haría falta para llevarlo a otro destino", eyebrow="Replicabilidad",
                subtitulo="El diseño es modular y replicable: es uno de los objetivos declarados del proyecto.")
cw2 = (ANCHO - Inches(0.5)) / 2
caja(s, MX, y, cw2, Inches(0.35), "LO QUE YA ESTÁ RESUELTO", size=11, bold=True, color=ACENTO)
vinetas(s, MX, y + Inches(0.42), cw2, Inches(3.3), [
    "**El método completo** y los modelos analíticos, ya validados sobre un destino real.",
    "**Las fuentes de ámbito nacional** —ACEVIN, Dataestur, Google Trends— que cubren cualquier ruta "
    "del vino española sin trabajo adicional.",
    "**El cuadro de mando**, con sus cinco vistas, su informe descargable y su guía metodológica.",
    "**La auditoría automática de webs** y el cálculo de accesibilidad y transporte, que funcionan "
    "con cualquier listado de empresas.",
], size=13.5, space=8)
caja(s, MX + cw2 + Inches(0.5), y, cw2, Inches(0.35), "LO QUE HABRÍA QUE APORTAR", size=11, bold=True, color=ACENTO)
vinetas(s, MX + cw2 + Inches(0.5), y + Inches(0.42), cw2, Inches(3.3), [
    "**El censo de empresas del destino** con su ubicación y servicios: el equivalente al catálogo de "
    "la Ruta del Vino y Brandy.",
    "**La descarga del corpus de reseñas** de ese destino, que es el único coste variable relevante.",
    "**Criterio local** para calibrar los umbrales de las recomendaciones y validar los resultados.",
    "Opcionalmente, **los datos que no existen en abierto** —empleo, consumos, precios— mediante "
    "cuestionario a las empresas.",
], size=13.5, space=8)
remate(s, "El proyecto está a mitad de sus 24 meses. Nada de lo que han visto es definitivo, y eso también es una oportunidad.",
       y=y + Inches(3.85))
notas(s, """
Honestidad sobre el estado: estamos a mitad del proyecto. Para un gestor interesado, eso significa
que todavía se puede influir en qué indicadores se priorizan.
No prometer plazos ni compromisos de despliegue en su destino: la conversación de hoy es de
conocimiento mutuo. Si hay interés real, se canaliza después.
""")

# ════════════════════════════ 12 · CIERRE ════════════════════════════
s = b.oscura("¿Hablamos?",
             eyebrow="Para seguir",
             bajada="La plataforma está abierta y se puede explorar hoy mismo. Si quieren verla sobre "
                    "su propio destino, o que se la enseñemos a sus empresas, hablemos.")
caja(s, MX, Inches(5.1), ANCHO, Inches(0.35), "LA PLATAFORMA", size=10, bold=True, color=rgb("EBD98A"))
caja(s, MX, Inches(5.4), ANCHO, Inches(0.4), "enolytics-feder-uca.streamlit.app", size=17, color=CREMA)
caja(s, MX, Inches(5.95), ANCHO, Inches(0.35), "CONTACTO", size=10, bold=True, color=rgb("EBD98A"))
caja(s, MX, Inches(6.25), ANCHO, Inches(0.5),
     "inteligencia.turistica@gm.uca.es · Universidad de Cádiz · Instituto INDESS\n"
     "Proyecto ENOLYTICS · FEDER Andalucía 2021-2027",
     size=13, color=CREMA)
notas(s, """
Preguntas probables:
· «¿Cuánto cuesta?» → el prototipo corre sobre software libre y servidor gratuito; el coste real es
  la actualización del corpus de reseñas, decenas de euros al año. El coste verdadero es el tiempo
  de quien lo mantiene y lo interpreta.
· «¿Puede ver una bodega los datos de otra?» → sí, y es una decisión consciente: sin comparación no
  hay inteligencia competitiva. Todo lo que se muestra es público.
· «¿Sirve para un destino que no sea de vino?» → el método sí; el léxico de atributos habría que
  reconstruirlo para ese sector.
· «¿Y la protección de datos?» → reseñas públicas, análisis agregados, el autor no interviene.
""")

# ════════════════════════════ GUARDAR ════════════════════════════
salida = pathlib.Path(sys.argv[1]) if len(sys.argv) > 1 else pathlib.Path(
    "docs/presentaciones/ENOLYTICS_gestores_de_destino.pptx")
ruta, n = b.guardar(salida)
print(f"guardado: {ruta} · {n} diapositivas")
